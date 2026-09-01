"""Leitura e criação local de DOCX sem alterar o arquivo importado."""

from __future__ import annotations

from dataclasses import asdict, dataclass
from io import BytesIO
from pathlib import PurePosixPath
from typing import Any
from zipfile import BadZipFile, ZipFile

from docx import Document

DOCX_MIME_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
MAX_INPUT_BYTES = 64 * 1024 * 1024
MAX_MEMBERS = 2_000
MAX_UNCOMPRESSED_BYTES = 256 * 1024 * 1024
MAX_COMPRESSION_RATIO = 200
MAX_PARAGRAPHS = 10_000
MAX_TABLES = 500
MAX_CELLS = 50_000
MAX_TEXT_LENGTH = 1_000_000
ALLOWED_STYLES = {
    "Normal",
    "Title",
    "Heading 1",
    "Heading 2",
    "Heading 3",
    "Quote",
    "List Bullet",
    "List Number",
}


class DocxInputError(ValueError):
    """Entrada DOCX rejeitada antes do processamento."""


@dataclass(frozen=True, slots=True)
class DocxParagraph:
    text: str
    style: str


@dataclass(frozen=True, slots=True)
class DocxInspection:
    paragraphs: tuple[DocxParagraph, ...]
    tables: tuple[tuple[tuple[str, ...], ...], ...]
    title: str | None

    def to_dict(self) -> dict[str, Any]:
        return {
            "paragraphs": [asdict(paragraph) for paragraph in self.paragraphs],
            "tables": [[list(row) for row in table] for table in self.tables],
            "title": self.title,
        }


def inspect_docx(content: bytes) -> DocxInspection:
    _validate_archive(content)
    try:
        document = Document(BytesIO(content))
    except (BadZipFile, KeyError, ValueError) as error:
        raise DocxInputError("O conteúdo não é um DOCX válido.") from error

    paragraphs = tuple(
        DocxParagraph(text=paragraph.text, style=paragraph.style.name or "Normal")
        for paragraph in document.paragraphs
    )
    tables = tuple(
        tuple(tuple(cell.text for cell in row.cells) for row in table.rows)
        for table in document.tables
    )
    return DocxInspection(
        paragraphs=paragraphs,
        tables=tables,
        title=document.core_properties.title or None,
    )


def create_docx(params: dict[str, Any]) -> bytes:
    title = _optional_text(params.get("title"), "title", 512)
    raw_paragraphs = params.get("paragraphs", [])
    raw_tables = params.get("tables", [])
    if not isinstance(raw_paragraphs, list) or len(raw_paragraphs) > MAX_PARAGRAPHS:
        raise DocxInputError("paragraphs deve ser uma lista com até 10.000 itens.")
    if not isinstance(raw_tables, list) or len(raw_tables) > MAX_TABLES:
        raise DocxInputError("tables deve ser uma lista com até 500 itens.")

    document = Document()
    if title:
        document.core_properties.title = title
    for index, item in enumerate(raw_paragraphs):
        if not isinstance(item, dict):
            raise DocxInputError(f"paragraphs[{index}] deve ser um objeto.")
        text = _required_text(item.get("text"), f"paragraphs[{index}].text")
        style = item.get("style", "Normal")
        if style not in ALLOWED_STYLES:
            raise DocxInputError(f"Estilo não permitido em paragraphs[{index}].")
        document.add_paragraph(text, style=style)

    cell_count = 0
    for table_index, raw_table in enumerate(raw_tables):
        if not isinstance(raw_table, list) or not raw_table:
            raise DocxInputError(f"tables[{table_index}] deve conter ao menos uma linha.")
        column_count: int | None = None
        normalized_rows: list[list[str]] = []
        for row_index, raw_row in enumerate(raw_table):
            if not isinstance(raw_row, list) or not raw_row:
                raise DocxInputError(
                    f"tables[{table_index}][{row_index}] deve conter ao menos uma célula."
                )
            if column_count is None:
                column_count = len(raw_row)
            elif len(raw_row) != column_count:
                raise DocxInputError(
                    f"Todas as linhas de tables[{table_index}] devem ter igual tamanho."
                )
            normalized_rows.append(
                [
                    _required_text(value, f"tables[{table_index}][{row_index}][{cell_index}]")
                    for cell_index, value in enumerate(raw_row)
                ]
            )
            cell_count += len(raw_row)
            if cell_count > MAX_CELLS:
                raise DocxInputError("As tabelas excedem o limite de 50.000 células.")
        table = document.add_table(rows=len(normalized_rows), cols=column_count or 1)
        table.style = "Table Grid"
        for row_index, row in enumerate(normalized_rows):
            for cell_index, value in enumerate(row):
                table.cell(row_index, cell_index).text = value

    output = BytesIO()
    document.save(output)
    content = output.getvalue()
    if len(content) > MAX_INPUT_BYTES:
        raise DocxInputError("O DOCX gerado excede o limite de 64 MiB.")
    return content


def _validate_archive(content: bytes) -> None:
    if not content or len(content) > MAX_INPUT_BYTES:
        raise DocxInputError("O arquivo deve ter entre 1 byte e 64 MiB.")
    try:
        with ZipFile(BytesIO(content)) as archive:
            members = archive.infolist()
            if len(members) > MAX_MEMBERS:
                raise DocxInputError("O DOCX excede o limite de 2.000 itens internos.")
            total_size = 0
            for member in members:
                path = PurePosixPath(member.filename.replace("\\", "/"))
                if path.is_absolute() or ".." in path.parts:
                    raise DocxInputError("O DOCX contém caminho interno inseguro.")
                unix_mode = member.external_attr >> 16
                if unix_mode & 0o170000 == 0o120000:
                    raise DocxInputError("O DOCX contém link simbólico não permitido.")
                if member.flag_bits & 0x1:
                    raise DocxInputError("DOCX criptografado não é suportado.")
                total_size += member.file_size
                if total_size > MAX_UNCOMPRESSED_BYTES:
                    raise DocxInputError("O DOCX excede o limite descompactado de 256 MiB.")
                if member.file_size > MAX_COMPRESSION_RATIO * max(member.compress_size, 1):
                    raise DocxInputError("O DOCX possui taxa de compressão insegura.")
            if (
                "[Content_Types].xml" not in archive.namelist()
                or "word/document.xml" not in archive.namelist()
            ):
                raise DocxInputError("O conteúdo não possui a estrutura mínima de um DOCX.")
    except BadZipFile as error:
        raise DocxInputError("O conteúdo não é um arquivo ZIP válido.") from error


def _required_text(value: object, field: str) -> str:
    if not isinstance(value, str) or len(value) > MAX_TEXT_LENGTH:
        raise DocxInputError(f"{field} deve ser texto com até 1.000.000 de caracteres.")
    return value


def _optional_text(value: object, field: str, limit: int) -> str | None:
    if value is None:
        return None
    if not isinstance(value, str) or len(value) > limit:
        raise DocxInputError(f"{field} deve ser texto com até {limit} caracteres.")
    return value
