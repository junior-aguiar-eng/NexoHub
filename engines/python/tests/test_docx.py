from __future__ import annotations

import base64
from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from docx import Document

from nexohub_document_engine.docx import DocxInputError, create_docx, inspect_docx
from nexohub_document_engine.protocol import handle_request


def test_create_and_inspect_docx_without_mutating_source() -> None:
    generated = create_docx(
        {
            "title": "Documento sintético",
            "paragraphs": [
                {"text": "NexoHub", "style": "Title"},
                {"text": "Conteúdo local.", "style": "Normal"},
            ],
            "tables": [[["Campo", "Valor"], ["Origem", "Teste"]]],
        }
    )
    snapshot = bytes(generated)

    inspection = inspect_docx(generated)

    assert generated == snapshot
    assert inspection.title == "Documento sintético"
    assert [paragraph.text for paragraph in inspection.paragraphs] == ["NexoHub", "Conteúdo local."]
    assert inspection.tables == ((("Campo", "Valor"), ("Origem", "Teste")),)
    reopened = Document(BytesIO(generated))
    assert reopened.paragraphs[0].text == "NexoHub"


def test_rejects_archive_with_path_traversal() -> None:
    output = BytesIO()
    with ZipFile(output, "w", ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", "<Types />")
        archive.writestr("word/document.xml", "<document />")
        archive.writestr("../escape.txt", "unsafe")

    with pytest.raises(DocxInputError, match="caminho interno inseguro"):
        inspect_docx(output.getvalue())


def test_protocol_creates_and_inspects_docx() -> None:
    created = handle_request(
        {
            "id": "create-1",
            "method": "docx.create",
            "params": {"paragraphs": [{"text": "Olá", "style": "Heading 1"}]},
        }
    )
    result = created["result"]
    assert result["mimeType"].endswith("wordprocessingml.document")
    assert result["sizeBytes"] > 0

    inspected = handle_request(
        {
            "id": "inspect-1",
            "method": "docx.inspect",
            "params": {"contentBase64": result["contentBase64"]},
        }
    )
    assert inspected["result"]["paragraphs"] == [{"text": "Olá", "style": "Heading 1"}]
    assert base64.b64decode(result["contentBase64"], validate=True)


def test_create_rejects_unknown_style() -> None:
    with pytest.raises(DocxInputError, match="Estilo não permitido"):
        create_docx({"paragraphs": [{"text": "Texto", "style": "External Style"}]})
