"""Benchmark descartável; não é runtime nem dependência do aplicativo.

Execute com o Python do venv isolado. Resultados e modelos ficam no diretório
ignorado .superpowers/sdd/2026-10-26-xl-c1/isolated.
"""

import argparse
import hashlib
import importlib.metadata
import json
import os
import platform
import queue
import subprocess
import sys
import threading
import time
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
CORPUS = ROOT / "apps/client/src/platform/__fixtures__/xl-c1"
WORK = ROOT / ".superpowers/sdd/2026-10-26-xl-c1/isolated"


def write_json(path, data):
    path.write_text(
        json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )


def score_cells(actual, expected):
    """Não corrige valores nem remove colunas. Alinha somente cabeçalho exato.

    Os extras são contabilizados separadamente. Ausência de tabela é ausência
    inclusive de células esperadas vazias, nunca acerto por preenchimento.
    """
    header = expected[0]
    offset = next((i for i, row in enumerate(actual) if row == header), 0)
    aligned = actual[offset:]
    cells = []
    strict_correct = 0
    for r, row in enumerate(expected):
        for c, value in enumerate(row):
            present = r < len(aligned) and c < len(aligned[r])
            found = aligned[r][c] if present else None
            strict_correct += int(
                r < len(actual) and c < len(actual[r]) and actual[r][c] == value
            )
            cells.append(
                {
                    "row": r + 1,
                    "column": c + 1,
                    "expected": value,
                    "actual": found,
                    "present": present,
                    "correct": present and found == value,
                }
            )
    correct = sum(cell["correct"] for cell in cells)
    extra = sum(len(row) for row in actual[:offset])
    extra += sum(
        len(row) if r >= len(expected) else max(0, len(row) - len(expected[r]))
        for r, row in enumerate(aligned)
    )
    return {
        "correct": correct,
        "expected": len(cells),
        "strictCorrect": strict_correct,
        "headerOffset": offset,
        "extraCells": extra,
        "cells": cells,
    }


def event(**values):
    print(json.dumps(values), flush=True)


def child(profile, pdf, output):
    start = time.perf_counter()
    event(phase="initialize")
    if profile.startswith("camelot"):
        import camelot
        import cv2

        cv2.setNumThreads(2)
        flavor = profile.split("-")[1]

        def extract():
            options = {"flavor": flavor, "pages": "all", "parallel": False}
            if profile.endswith("roi"):
                options.update(
                    pages="2", table_areas=["0,720,600,560"], columns=["180,372"]
                )
            tables = camelot.read_pdf(str(pdf), **options)
            return [
                {
                    "page": int(t.page),
                    "rows": t.df.values.tolist(),
                    "bbox": list(t._bbox),
                    "report": t.parsing_report,
                }
                for t in tables
            ]
    else:
        from docling.datamodel.accelerator_options import (
            AcceleratorDevice,
            AcceleratorOptions,
        )
        from docling.datamodel.base_models import InputFormat
        from docling.datamodel.pipeline_options import (
            PdfPipelineOptions,
            TableFormerMode,
        )
        from docling.document_converter import DocumentConverter, PdfFormatOption

        options = PdfPipelineOptions(
            artifacts_path=WORK / "models",
            do_ocr=False,
            enable_remote_services=False,
            do_table_structure=True,
        )
        options.table_structure_options.mode = TableFormerMode.ACCURATE
        options.table_structure_options.do_cell_matching = True
        options.accelerator_options = AcceleratorOptions(
            num_threads=2, device=AcceleratorDevice.CPU
        )
        converter = DocumentConverter(
            format_options={InputFormat.PDF: PdfFormatOption(pipeline_options=options)}
        )
        converter.initialize_pipeline(InputFormat.PDF)

        def extract():
            result = converter.convert(pdf, raises_on_error=True)
            tables = []
            for table in result.document.tables:
                data = table.data
                rows = [[""] * data.num_cols for _ in range(data.num_rows)]
                # Âncora do span; não duplica texto de célula mesclada.
                for cell in data.table_cells:
                    rows[cell.start_row_offset_idx][cell.start_col_offset_idx] = (
                        cell.text
                    )
                tables.append(
                    {
                        "page": table.prov[0].page_no,
                        "rows": rows,
                        "bbox": table.prov[0].bbox.model_dump(mode="json"),
                        "cells": [
                            cell.model_dump(mode="json") for cell in data.table_cells
                        ],
                    }
                )
            return tables

    initialization_ms = (time.perf_counter() - start) * 1000
    runs = []
    for index in range(3):
        event(phase=f"extract-{index + 1}")
        before = time.perf_counter()
        tables = extract()
        elapsed = (time.perf_counter() - before) * 1000
        event(phase="between")
        runs.append({"pass": index + 1, "extractionMs": elapsed, "tables": tables})
    write_json(
        output,
        {"profile": profile, "initializationMs": initialization_ms, "runs": runs},
    )


def observe(command, prefix, env):
    import psutil

    phases = {}
    current = "startup"
    peak = 0
    os_peak = 0
    events = queue.Queue()
    started = time.perf_counter()
    with prefix.with_suffix(".stderr.log").open("w", encoding="utf-8") as errors:
        process = subprocess.Popen(
            command,
            stdout=subprocess.PIPE,
            stderr=errors,
            env=env,
            encoding="utf-8",
            errors="replace",
        )

        def read_events():
            for line in process.stdout:
                try:
                    events.put(json.loads(line))
                except json.JSONDecodeError:
                    pass

        thread = threading.Thread(target=read_events, daemon=True)
        thread.start()
        parent = psutil.Process(process.pid)
        while process.poll() is None:
            while not events.empty():
                entry = events.get_nowait()
                current = entry.get("phase", current)
            rss = 0
            try:
                descendants = [parent, *parent.children(recursive=True)]
            except psutil.Error:
                descendants = []
            for member in descendants:
                try:
                    info = member.memory_info()
                    rss += info.rss
                    if member.pid == parent.pid:
                        os_peak = max(os_peak, getattr(info, "peak_wset", 0))
                except psutil.Error:
                    pass
            peak = max(peak, rss)
            phases[current] = max(phases.get(current, 0), rss)
            if time.perf_counter() - started > 600:
                for member in reversed(descendants):
                    try:
                        member.kill()
                    except psutil.Error:
                        pass
                raise TimeoutError("Caso excedeu 600 segundos")
            time.sleep(0.02)
        thread.join(timeout=2)
    if process.returncode:
        raise RuntimeError(
            f"Caso falhou ({process.returncode}); ver {prefix.with_suffix('.stderr.log')}"
        )
    return {
        "processWallMs": (time.perf_counter() - started) * 1000,
        "sampledTreePeakRssBytes": peak,
        "windowsLaunchedProcessPeakWorkingSetBytes": os_peak,
        "sampledPhaseRssBytes": phases,
        "sampleIntervalMs": 20,
    }


def evaluate(case, manifest):
    expected_pages = manifest["expectedCellsByPage"]
    pages = [2] if case["profile"].endswith("roi") else [1, 2]
    for run in case["runs"]:
        scores = []
        for page in pages:
            # Ordem original; nenhum candidato é escolhido por se aproximar do gabarito.
            candidates = [t for t in run["tables"] if t["page"] == page]
            actual = candidates[0]["rows"] if candidates else []
            score = score_cells(actual, expected_pages[page - 1])
            score.update(
                page=page,
                detectedTables=len(candidates),
                extraTables=max(0, len(candidates) - 1),
            )
            scores.append(score)
        run["quality"] = scores
        run["correctCells"] = sum(s["correct"] for s in scores)
        run["expectedCells"] = sum(s["expected"] for s in scores)
        run["extraCells"] = sum(s["extraCells"] for s in scores)


def export_and_verify(case, path):
    from openpyxl import Workbook, load_workbook

    workbook = Workbook()
    workbook.remove(workbook.active)
    for i, table in enumerate(case["runs"][0]["tables"]):
        sheet = workbook.create_sheet(f"Pagina{table['page']}-Tabela{i + 1}")
        for r, row in enumerate(table["rows"], 1):
            for c, value in enumerate(row, 1):
                cell = sheet.cell(r, c, value)
                cell.data_type = (
                    "s"  # Conteúdo literal, sem interpretar códigos/fórmulas.
                )
    if not workbook.worksheets:
        workbook.create_sheet("Sem tabelas")
    workbook.save(path)
    reopened = load_workbook(path)
    for i, table in enumerate(case["runs"][0]["tables"]):
        sheet = reopened.worksheets[i]
        for r, row in enumerate(table["rows"], 1):
            for c, value in enumerate(row, 1):
                assert sheet.cell(r, c).value == (value if value != "" else None)
                if value:
                    assert sheet.cell(r, c).data_type == "s"
    return {
        "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
        "bytes": path.stat().st_size,
        "reopenedCellValuesVerified": True,
        "exporter": "openpyxl compartilhado, fora do tempo de extração; sem integração",
    }


def main():
    WORK.mkdir(parents=True, exist_ok=True)
    manifest = json.loads((CORPUS / "manifest.json").read_text(encoding="utf-8"))
    env = dict(
        os.environ,
        PYTHONUTF8="1",
        PYTHONIOENCODING="utf-8",
        HF_HUB_OFFLINE="1",
        HF_HUB_DISABLE_TELEMETRY="1",
        DO_NOT_TRACK="1",
        TRANSFORMERS_OFFLINE="1",
        HF_HOME=str(WORK / "hf-cache"),
        OMP_NUM_THREADS="2",
        MKL_NUM_THREADS="2",
        OPENBLAS_NUM_THREADS="2",
        HF_HUB_DISABLE_XET="1",
    )
    results = []
    for profile in [
        "camelot-lattice",
        "camelot-stream",
        "camelot-network",
        "camelot-stream-roi",
        "docling-accurate",
        "pdfjs-auto",
        "pdfjs-roi",
    ]:
        for fixture in manifest["fixtures"]:
            pdf = CORPUS / fixture["name"]
            before = hashlib.sha256(pdf.read_bytes()).hexdigest()
            assert before == fixture["sha256"]
            label = f"{profile}-{pdf.stem}"
            output = WORK / f"{label}.json"
            command = [
                sys.executable,
                str(Path(__file__).resolve()),
                "--child",
                profile,
                "--pdf",
                str(pdf),
                "--output",
                str(output),
            ]
            if profile.startswith("pdfjs"):
                command = [
                    "node",
                    str(ROOT / "scripts/comparar-xl-c1-browser.mjs"),
                    profile,
                    str(pdf),
                    str(output),
                ]
            observation = observe(command, WORK / label, env)
            case = json.loads(output.read_text(encoding="utf-8"))
            case.update(
                fixture=fixture["name"],
                sourceSha256=before,
                sourceUnchanged=hashlib.sha256(pdf.read_bytes()).hexdigest() == before,
                resources=observation,
            )
            assert case["sourceUnchanged"]
            evaluate(case, manifest)
            case["export"] = export_and_verify(case, WORK / f"{label}.xlsx")
            results.append(case)
            write_json(WORK / "comparison.json", {"cases": results})
            print(
                label,
                [
                    (r["correctCells"], r["expectedCells"], round(r["extractionMs"], 1))
                    for r in case["runs"]
                ],
                flush=True,
            )
    packages = {
        d.metadata["Name"]: d.version for d in importlib.metadata.distributions()
    }
    write_json(
        WORK / "comparison.json",
        {
            "measuredAt": time.strftime("%Y-%m-%dT%H:%M:%SZ", time.gmtime()),
            "python": sys.version,
            "platform": platform.platform(),
            "cpu": platform.processor(),
            "packages": packages,
            "manifest": manifest,
            "cases": results,
            "method": "CPU 2 threads. Processo novo por perfil/PDF; 3 extrações no mesmo processo. Import/modelos separados de extração; sem download durante medição. RSS da árvore amostrado em 20ms, pico working set Windows do processo. Qualidade exata sem corrigir valores; cabeçalho exato pode alinhar linhas, extras explícitos. Primeiro candidato por página, sem seleção por qualidade. XLSX compartilhado fora da medição.",
        },
    )


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--child")
    parser.add_argument("--pdf", type=Path)
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()
    if args.child:
        child(args.child, args.pdf, args.output)
    else:
        main()
